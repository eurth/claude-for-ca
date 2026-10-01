"""Excel registers and the 11-sheet Master Accounts workbook."""

from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter


HEADER = Font(bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="1F5C4D")


def _num(v) -> float:
    try:
        return float(v or 0)
    except (TypeError, ValueError):
        return 0.0


def _write_sheet(wb: Workbook, title: str, headers: list[str], rows: list[list[Any]]) -> None:
    ws = wb.create_sheet(title)
    for col, head in enumerate(headers, start=1):
        cell = ws.cell(1, col, head)
        cell.font = HEADER
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(wrap_text=True)
    for r_i, row in enumerate(rows, start=2):
        for c_i, val in enumerate(row, start=1):
            ws.cell(r_i, c_i, val)
    for i in range(1, len(headers) + 1):
        ws.column_dimensions[get_column_letter(i)].width = 16
    if "Sheet" in wb.sheetnames and len(wb.sheetnames) > 1:
        del wb["Sheet"]


def write_rows_xlsx(path: Path, sheet: str, headers: list[str], rows: list[list[Any]]) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    _write_sheet(wb, sheet, headers, rows)
    if "Sheet" in wb.sheetnames:
        del wb["Sheet"]
    wb.save(path)
    return path


def purchase_rows_table(rows: list[dict]) -> tuple[list[str], list[list[Any]]]:
    headers = ["Vendor", "GSTIN", "Invoice", "Date", "HSN", "Description", "Qty", "Taxable", "CGST", "SGST", "IGST", "Total"]
    out = []
    for row in rows:
        out.append(
            [
                row.get("vendor") or row.get("vendor_name"),
                row.get("gstin") or row.get("vendor_gstin"),
                row.get("invoice_no") or row.get("invoice_number"),
                row.get("date") or row.get("invoice_date"),
                row.get("hsn"),
                row.get("description"),
                row.get("qty"),
                _num(row.get("taxable_value") or row.get("taxable")),
                _num(row.get("cgst") or row.get("cgst_amount")),
                _num(row.get("sgst") or row.get("sgst_amount")),
                _num(row.get("igst") or row.get("igst_amount")),
                _num(row.get("total")),
            ]
        )
    return headers, out


def bank_rows_table(rows: list[dict]) -> tuple[list[str], list[list[Any]]]:
    headers = ["Date", "Narration", "Debit", "Credit", "Balance", "Category"]
    out = [
        [r.get("date"), r.get("narration"), _num(r.get("debit")), _num(r.get("credit")), _num(r.get("balance")), r.get("category")]
        for r in rows
    ]
    return headers, out


def json_table_xlsx(path: Path, title: str, data: dict | None) -> Path:
    data = data or {}
    rows = data.get("rows")
    if isinstance(rows, list) and rows and isinstance(rows[0], dict):
        headers = list(rows[0].keys())
        body = [[row.get(h) for h in headers] for row in rows]
        return write_rows_xlsx(path, title[:31], headers, body)
    items = data.get("items") or data.get("flags")
    if isinstance(items, list):
        return write_rows_xlsx(path, title[:31], ["Item"], [[str(i)] for i in items])
    headers = ["Field", "Value"]
    body = [[k, str(v)] for k, v in data.items() if k not in ("rows", "items", "flags")]
    return write_rows_xlsx(path, title[:31] or "Working", headers, body or [["(empty)", ""]])


def build_master_accounts(
    path: Path,
    *,
    company: str,
    period: str | None,
    purchase: list[dict],
    sales: list[dict],
    bank: list[dict],
    opening: float | None = None,
    closing: float | None = None,
) -> dict[str, Any]:
    path.parent.mkdir(parents=True, exist_ok=True)
    taxable = sum(_num(r.get("taxable_value") or r.get("taxable")) for r in purchase)
    cgst = sum(_num(r.get("cgst") or r.get("cgst_amount")) for r in purchase)
    sgst = sum(_num(r.get("sgst") or r.get("sgst_amount")) for r in purchase)
    igst = sum(_num(r.get("igst") or r.get("igst_amount")) for r in purchase)
    purchase_total = sum(_num(r.get("total")) for r in purchase)
    sales_total = sum(_num(r.get("total")) for r in sales)
    sales_taxable = sum(_num(r.get("taxable_value") or r.get("taxable")) for r in sales)
    bank_debit = sum(_num(r.get("debit")) for r in bank)
    bank_credit = sum(_num(r.get("credit")) for r in bank)
    gst_pay = max(sales_taxable * 0.18 - (cgst + sgst + igst), 0) if sales else max(0, -(cgst + sgst + igst) + 0)

    wb = Workbook()
    dash = wb.active
    dash.title = "SUMMARY"
    dash["A1"] = company
    dash["A1"].font = Font(bold=True, size=14)
    dash["A2"] = f"Period: {period or '—'}"
    dash["A4"] = "Purchases (taxable)"
    dash["B4"] = taxable
    dash["A5"] = "Purchases (gross)"
    dash["B5"] = purchase_total
    dash["A6"] = "Sales (gross)"
    dash["B6"] = sales_total
    dash["A7"] = "ITC CGST+SGST+IGST"
    dash["B7"] = cgst + sgst + igst
    dash["A8"] = "Bank debit"
    dash["B8"] = bank_debit
    dash["A9"] = "Bank credit"
    dash["B9"] = bank_credit
    dash["A10"] = "Opening"
    dash["B10"] = opening
    dash["A11"] = "Closing"
    dash["B11"] = closing
    dash["A13"] = "GST payable (indicative from books — not a filed 3B)"
    dash["B13"] = gst_pay

    p_h, p_rows = purchase_rows_table(purchase)
    _write_sheet(wb, "PURCHASE", p_h, p_rows)
    s_h, s_rows = purchase_rows_table(sales)
    _write_sheet(wb, "SALES", s_h, s_rows)
    b_h, b_rows = bank_rows_table(bank)
    _write_sheet(wb, "BANK", b_h, b_rows)
    cash = [r for r in bank if "CASH" in str(r.get("narration") or "").upper()]
    _write_sheet(wb, "CASH", b_h, bank_rows_table(cash)[1])
    _write_sheet(
        wb,
        "GST",
        ["Head", "Amount"],
        [
            ["Outward taxable (sales)", sales_taxable],
            ["ITC CGST", cgst],
            ["ITC SGST", sgst],
            ["ITC IGST", igst],
            ["Indicative net GST", gst_pay],
        ],
    )
    tds_rows = [
        [r.get("date"), r.get("vendor") or r.get("narration"), r.get("gstin"), "194C?", r.get("total") or r.get("debit")]
        for r in purchase + bank
        if "TDS" in str(r.get("narration") or r.get("description") or "").upper()
    ]
    _write_sheet(wb, "TDS", ["Date", "Party", "GSTIN/PAN", "Section", "Amount"], tds_rows)
    _write_sheet(
        wb,
        "PAYABLES",
        ["Vendor", "Invoice", "Total", "Age bucket"],
        [[r.get("vendor"), r.get("invoice_no"), _num(r.get("total")), "Unpaid (books)"] for r in purchase],
    )
    _write_sheet(
        wb,
        "RECEIVABLES",
        ["Customer", "Invoice", "Total", "Age bucket"],
        [[r.get("vendor"), r.get("invoice_no"), _num(r.get("total")), "Unpaid (books)"] for r in sales],
    )
    _write_sheet(
        wb,
        "P&L",
        ["Head", period or "YTD"],
        [
            ["Revenue from operations", sales_taxable],
            ["Cost of goods / purchases", taxable],
            ["Gross profit", sales_taxable - taxable],
        ],
    )
    _write_sheet(
        wb,
        "BS",
        ["Head", "Amount"],
        [["Bank (closing)", closing], ["Payables (purchases)", purchase_total]],
    )
    wb.save(path)
    return {
        "path": str(path),
        "purchase_rows": len(purchase),
        "sales_rows": len(sales),
        "bank_rows": len(bank),
        "purchase_total": purchase_total,
        "sales_total": sales_total,
        "itc": cgst + sgst + igst,
    }
